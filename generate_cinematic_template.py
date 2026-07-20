#!/usr/bin/env python3
"""Regenerate cinematic-dark template.docx to match the high-quality PDF design."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from pathlib import Path

OUT = Path.home() / "athamu_workspace" / "templates" / "cinematic-dark"
OUT.mkdir(parents=True, exist_ok=True)

PAGE_W = Cm(21.0)
PAGE_H = Cm(29.7)
MARGIN = Cm(1.6)

DARK_BG = "0D0D11"
CARD_BG = "16161F"
GOLD = "FFB800"
RED = "E50914"
WHITE = "FFFFFF"
LIGHT_GRAY = "CCCCCC"


def rgb(color: str) -> RGBColor:
    color = color.lstrip("#")
    return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))


def set_cell_bg(cell, color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color.replace("#", ""))
    tcPr.append(shd)


def set_cell_margins(cell, top=Cm(0.25), bottom=Cm(0.25), left=Cm(0.4), right=Cm(0.4)):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [
        ("top", top),
        ("bottom", bottom),
        ("left", left),
        ("right", right),
    ]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(int(round(val.cm * 567))))  # 1 cm ≈ 567 twips
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_table_border(table, color: str = GOLD, size_pt: int = 1):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size_pt * 8))  # 1pt ≈ 8 eighths of a point
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color.replace("#", ""))
        borders.append(b)
    tblPr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)


def add_paragraph_style(doc, name, font_name, font_size, color, bold=False, italic=False, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=Pt(8)):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font_name
    style.font.size = Pt(font_size)
    style.font.color.rgb = rgb(color)
    style.font.bold = bold
    style.font.italic = italic
    style.paragraph_format.alignment = alignment
    style.paragraph_format.space_after = space_after
    return style


def set_page_background(doc, color: str):
    # Set background color for the document body.
    try:
        from docx.oxml.ns import nsmap as _nsmap
        body = doc.element.body
        bg = OxmlElement("w:background")
        bg.set(qn("w:color"), color.replace("#", ""))
        body.insert(0, bg)
    except Exception:
        pass


def create_decorative_line():
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GOLD)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.append(pPr)
    return p


def add_section_title(cell, text, color=GOLD):
    cell.paragraphs[0].clear()
    p = cell.paragraphs[0]
    p.style = doc.styles["SectionTitle"]
    run = p.add_run(text)
    run.font.color.rgb = rgb(color)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


doc = Document()
sec = doc.sections[0]
sec.page_height = PAGE_H
sec.page_width = PAGE_W
sec.top_margin = MARGIN
sec.bottom_margin = MARGIN
sec.left_margin = MARGIN
sec.right_margin = MARGIN

set_page_background(doc, DARK_BG)

# Base paragraph style
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.font.color.rgb = rgb(WHITE)
normal.paragraph_format.line_spacing = 1.15

# Structured styles
add_paragraph_style(doc, "Cinematic_Title", "Georgia", 28, GOLD, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(2))
add_paragraph_style(doc, "Cinematic_Subtitle", "Georgia", 18, LIGHT_GRAY, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(6))
add_paragraph_style(doc, "Cinematic_Body", "Times New Roman", 12, WHITE, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=Pt(8))
add_paragraph_style(doc, "SectionTitle", "Georgia", 16, GOLD, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=Pt(10))
add_paragraph_style(doc, "Cinematic_Caption", "Arial", 10, LIGHT_GRAY, italic=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(4))
add_paragraph_style(doc, "Cinematic_Footer", "Arial", 9, LIGHT_GRAY, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=Pt(0))

# Cover frame
cover = doc.add_table(rows=1, cols=1)
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
add_table_border(cover, GOLD, 3)
c = cover.cell(0, 0)
set_cell_bg(c, DARK_BG)
c.width = Cm(17.8)
c.height = Cm(26.0)
cp = c.paragraphs[0]
cp.style = doc.styles["Cinematic_Title"]
cp.text = "{{titulo}}"

c.add_paragraph("")
cp = c.add_paragraph("{{subtitulo}}")
cp.style = doc.styles["Cinematic_Subtitle"]

c.add_paragraph("")
cp = c.add_paragraph("{{descripcion_corta}}")
cp.style = doc.styles["Cinematic_Body"]
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

c.add_paragraph("")
cp = c.add_paragraph("ATHA Producciones  |  {{fecha}}")
cp.style = doc.styles["Cinematic_Caption"]

doc.add_paragraph("")


# Sinopsis section
sinopsis = doc.add_table(rows=1, cols=1)
add_table_border(sinopsis, GOLD, 2)
set_cell_bg(sinopsis.cell(0, 0), CARD_BG)
add_section_title(sinopsis.cell(0, 0), "SINOPSIS", GOLD)

p = sinopsis.cell(0, 0).add_paragraph("{{sinopsis}}")
p.style = doc.styles["Cinematic_Body"]
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

sinopsis.cell(0, 0).add_paragraph("")
cp = sinopsis.cell(0, 0).add_paragraph("[IMAGEN: {{imagen_1}}]")
cp.style = doc.styles["Cinematic_Caption"]
doc.add_paragraph("")


# Propuesta section
propuesta = doc.add_table(rows=1, cols=1)
add_table_border(propuesta, RED, 2)
set_cell_bg(propuesta.cell(0, 0), CARD_BG)
add_section_title(propuesta.cell(0, 0), "PROPUESTA ARTÍSTICA", RED)
p = propuesta.cell(0, 0).add_paragraph("{{propuesta_texto}}")
p.style = doc.styles["Cinematic_Body"]
doc.add_paragraph("")


# Requerimientos section
req = doc.add_table(rows=1, cols=1)
add_table_border(req, GOLD, 2)
set_cell_bg(req.cell(0, 0), CARD_BG)
add_section_title(req.cell(0, 0), "REQUERIMIENTOS TÉCNICOS", GOLD)
p = req.cell(0, 0).add_paragraph("{{requerimientos_texto}}")
p.style = doc.styles["Cinematic_Body"]
p = req.cell(0, 0).add_paragraph("{{requerimientos_detalle}}")
p.style = doc.styles["Cinematic_Body"]
doc.add_paragraph("")


# Contacto section
contacto = doc.add_table(rows=1, cols=1)
add_table_border(contacto, GOLD, 2)
set_cell_bg(contacto.cell(0, 0), CARD_BG)
add_section_title(contacto.cell(0, 0), "CONTACTO", GOLD)
for text in [
    "{{contacto_texto}}",
    "Email: {{email}}",
    "Web: {{web}}",
    "Redes: {{redes}}",
]:
    p = contacto.cell(0, 0).add_paragraph(text)
    p.style = doc.styles["Cinematic_Body"]

doc.add_paragraph("")

# Footer logo
fp = doc.add_paragraph("[LOGO ATHA: {{logo_atha}}]")
fp.style = doc.styles["Cinematic_Footer"]
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT


out = OUT / "template.docx"
doc.save(str(out))
print(out)
