#!/usr/bin/env python3
"""ATHAMU Fase 3+: pipeline híbrido DOCX editable + PDF cinematográfico desde HTML+CSS."""
import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from docxtpl import DocxTemplate
except Exception as e:
    print("Falta dependencia docxtpl:", e, file=sys.stderr)
    sys.exit(1)

try:
    from google.oauth2.credentials import Credentials
    from google.auth.transport.requests import Request
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
except Exception as e:
    print("Falta dependencia google-auth / google-api-python-client:", e, file=sys.stderr)
    sys.exit(1)

from docx import Document

WORKSPACE = Path.home() / "athamu_workspace"
DEFAULT_VARIANT = "cinematic-dark"
TEMPLATES_DIR = WORKSPACE / "templates"
TOKENS_PATH = WORKSPACE / "tokens" / "atha-design-tokens.json"
OUTPUT_DIR = WORKSPACE / "output"
TOKEN_PATH = Path.home() / ".config" / "athamu" / "google_slides_pipeline" / "token.json"

MARKERS = [
    "{{titulo}}", "{{artista}}", "{{subtitulo}}", "{{descripcion_corta}}", "{{fecha}}",
    "{{sinopsis}}", "{{propuesta_texto}}", "{{nota_direccion}}",
    "{{requerimientos_texto}}", "{{requerimientos_detalle}}",
    "{{contacto_texto}}", "{{cierre_texto}}",
    "{{imagen_1}}", "{{imagen_2}}",
    "{{email}}", "{{web}}", "{{redes}}",
    "{{logo_atha}}",
]
PLACEHOLDER_IMAGE = "https://drive.google.com/thumbnail?id=10CPr4Ywd2VnPl3GIfCBwK6ihflS6MS3n&sz=w2000"

SIMPLE_TEMPLATE_CANDIDATES = [
    WORKSPACE / "templates" / "docx-editable" / "template.docx",
    WORKSPACE / "templates" / DEFAULT_VARIANT / "template.docx",
]


def log_json(path: Path, data: dict):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def warn(warnings: list, message: str):
    warnings.append(str(message))


def normalize_image(value: Any) -> str:
    value = (value or "").strip() if value is not None else ""
    if not value:
        return ""
    if "drive.google.com" in value:
        try:
            file_id = value.split("id=")[-1].split("&")[0]
            return f"https://drive.google.com/thumbnail?id={file_id}&sz=w2000"
        except Exception:
            return PLACEHOLDER_IMAGE
    return str(value)


def load_payload(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Payload must be a JSON object")
    return data


def map_context(data: dict) -> dict:
    ctx = {}
    for marker in MARKERS:
        key = re.sub(r"\{\{|\}\}", "", marker).strip()
        ctx[key] = data.get(key, "")
    # normalize images
    ctx["imagen_1"] = normalize_image(ctx.get("imagen_1"))
    ctx["imagen_2"] = normalize_image(ctx.get("imagen_2"))
    ctx["logo_atha"] = normalize_image(ctx.get("logo_atha", ""))
    if not ctx.get("titulo"):
        ctx["titulo"] = "Dossier"
    return ctx


def extract_text_from_docx(docx_path: Path) -> dict:
    """Extrae texto y marcadores desde un DOCX editado para reconstruir el payload."""
    doc = Document(str(docx_path))
    data = {}
    current_key = None
    buffer = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Detectar marcadores de campo comunes en el DOCX editable
        marker_match = re.match(r"^\{\{(\w+)\}\}$", text, re.IGNORECASE)
        if marker_match:
            if current_key and buffer:
                data[current_key] = "\n".join(buffer).strip()
                buffer = []
            current_key = marker_match.group(1).lower()
            continue
        # Detectar líneas con prefijos conocidos
        prefixed = re.match(r"^(Email|Web|Redes|Nota de Dirección|Propuesta)\s*:(.*)$", text, re.IGNORECASE)
        if prefixed:
            if current_key and buffer:
                data[current_key] = "\n".join(buffer).strip()
                buffer = []
            key = prefixed.group(1).lower().replace(" ", "_")
            value = prefixed.group(2).strip()
            current_key = key
            if value:
                buffer.append(value)
            continue
        if current_key:
            buffer.append(text)

    if current_key and buffer:
        data[current_key] = "\n".join(buffer).strip()

    return data


def simple_docx_template_path() -> Path:
    for p in SIMPLE_TEMPLATE_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError(f"No DOCX template found: {SIMPLE_TEMPLATE_CANDIDATES}")


def render_simple_docx(tpl_path: Path, context: dict, out_path: Path, warnings: list) -> Path:
    temp_copy = None
    try:
        doc = Document(str(tpl_path))
        try:
            sec = doc.sections[0]
            if sec.page_width is None or sec.page_height is None:
                sec.page_height = Cm(29.7)
                sec.page_width = Cm(21.0)
        except Exception:
            pass
        import tempfile
        fd, temp_copy_path = tempfile.mkstemp(suffix=".docx")
        import os
        os.close(fd)
        doc.save(temp_copy_path)
        render_src = Path(temp_copy_path)
        temp_copy = temp_copy_path
    except Exception as e:
        warn(warnings, f"Normalización previa falló, usar plantilla original: {e}")
        render_src = tpl_path

    tpl = DocxTemplate(str(render_src))
    tpl.render(context)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tpl.save(str(out_path))

    if temp_copy:
        try:
            Path(temp_copy).unlink(missing_ok=True)
        except Exception:
            pass
    return out_path


def convert_to_pdf_libreoffice(docx_path: Path, out_dir: Path, warnings: list) -> Path | None:
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to", "pdf:writer_pdf_Export",
        "--outdir", str(out_dir),
        str(docx_path),
    ]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        if proc.returncode != 0:
            warnings.append(f"LibreOffice falló: rc={proc.returncode}, {proc.stderr}")
            return None
        expected = out_dir / f"{docx_path.stem}.pdf"
        if expected.exists():
            return expected
        warnings.append("LibreOffice: PDF de salida no encontrado")
        return None
    except FileNotFoundError:
        warnings.append("LibreOffice no disponible en PATH")
        return None
    except subprocess.TimeoutExpired:
        warnings.append("LibreOffice excedió timeout")
        return None


def render_html_dossier(data: dict, variant: str = DEFAULT_VARIANT) -> Path:
    """Llama a render_dossier.py en modo HTML para obtener el PDF cinematográfico."""
    import importlib.util
    render_spec = importlib.util.spec_from_file_location("render_dossier_module", str(WORKSPACE / "render_dossier.py"))
    render_mod = importlib.util.module_from_spec(render_spec)
    render_spec.loader.exec_module(render_mod)
    result = render_mod.render_dossier(data, variant)
    return Path(result["pdf"])


def drive_service():
    if not TOKEN_PATH.exists():
        raise FileNotFoundError(f"Token no encontrado: {TOKEN_PATH}")
    info = json.loads(TOKEN_PATH.read_text(encoding="utf-8"))
    scopes = [
        "https://www.googleapis.com/auth/drive",
        "https://www.googleapis.com/auth/presentations",
    ]
    creds = Credentials.from_authorized_user_info(info, scopes)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
    return build("drive", "v3", credentials=creds)


def upload_to_drive(path: Path, mime_type: str, drive_links: list, warnings: list):
    service = drive_service()
    file_metadata = {"name": path.name}
    media = MediaFileUpload(str(path), mimetype=mime_type, resumable=True)
    try:
        request = service.files().create(body=file_metadata, media_body=media, fields="id,name,webViewLink,webContentLink")
        response = request.execute()
        link = {
            "name": response.get("name"),
            "id": response.get("id"),
            "webViewLink": response.get("webViewLink"),
            "webContentLink": response.get("webContentLink"),
            "mimeType": mime_type,
        }
        drive_links.append(link)
        return link
    except Exception as e:
        warn(warnings, f"Upload Drive falló para {path}: {e}")
        return None


def process_payload(payload_path: Path, variant: str = DEFAULT_VARIANT, manual_docx: Path | None = None, pdf_only: bool = False) -> dict:
    warnings: list = []
    start = datetime.now(timezone.utc).isoformat()
    summary = {
        "timestamp": start,
        "payload_path": str(payload_path),
        "variant": variant,
        "template_used": None,
        "docx_path": None,
        "pdf_path": None,
        "drive_links": [],
        "warnings": warnings,
        "success": False,
        "notes": "",
    }

    payload = load_payload(payload_path)
    context = map_context(payload)
    output_title = context.get("titulo") or payload.get("titulo") or "dossier"
    safe_title = re.sub(r"[^A-Za-z0-9_\-]+", "_", str(output_title)).strip("_") or "dossier"
    out_dir = OUTPUT_DIR / variant / safe_title
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1) DOCX editable
    if not pdf_only:
        try:
            tpl_path = simple_docx_template_path()
            summary["template_used"] = str(tpl_path)
            docx_output = out_dir / f"{safe_title}__{variant}_editable.docx"
            render_simple_docx(tpl_path, context, docx_output, warnings)
            summary["docx_path"] = str(docx_output)
        except Exception as e:
            warn(warnings, f"DOCX render failed: {e}")
    else:
        docx_output = None

    # Si viene un DOCX editado, extraer datos para regenerar HTML/PDF
    if manual_docx:
        if not manual_docx.exists():
            warn(warnings, f"--manual-docx no existe: {manual_docx}")
            summary["notes"] = "manual_docx missing"
            return summary
        extracted = extract_text_from_docx(manual_docx)
        if extracted:
            payload.update(extracted)
            context = map_context(payload)

    # 2) PDF desde HTML+CSS (autoridad visual)
    try:
        pdf_path = render_html_dossier(payload, variant)
        summary["pdf_path"] = str(pdf_path)
    except Exception as e:
        warn(warnings, f"PDF desde HTML falló: {e}")

    # 3) Subir a Drive
    drive_links = []
    for p, mime in [
        (Path(summary.get("pdf_path") if summary.get("pdf_path") else ""), "application/pdf"),
        (Path(summary.get("docx_path") if summary.get("docx_path") else ""), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ]:
        if p and p.exists():
            upload_to_drive(p, mime, drive_links, warnings)
    summary["drive_links"] = drive_links
    summary["success"] = bool(summary.get("pdf_path")) or bool(drive_links)

    log_path = out_dir / f"{safe_title}_generate_log.json"
    log_json(log_path, summary)
    return summary


def main():
    import subprocess
    parser = argparse.ArgumentParser(description="ATHAMU pipeline híbrido: DOCX editable + PDF cinematográfico")
    parser.add_argument("payload", type=str, help="Path to payload JSON")
    parser.add_argument("variant", nargs="?", default=DEFAULT_VARIANT, help="Template variant: cinematic-dark | light")
    parser.add_argument("--manual-docx", type=str, default=None, help="Regenerar PDF desde DOCX editado")
    parser.add_argument("--pdf-only", action="store_true", help="Solo PDF desde HTML (sin DOCX)")
    args = parser.parse_args()

    payload_path = Path(args.payload)
    if not payload_path.exists():
        print(json.dumps({"error": f"Payload no encontrado: {payload_path}"}, ensure_ascii=False))
        sys.exit(2)

    summary = process_payload(payload_path, args.variant, Path(args.manual_docx) if args.manual_docx else None, args.pdf_only)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
