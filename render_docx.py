#!/usr/bin/env python3
"""Generate editable DOCX from dossier data with matching markers."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

WORKSPACE = Path.home() / "athamu_workspace"
PLACEHOLDER = "https://drive.google.com/thumbnail?id=10CPr4Ywd2VnPl3GIfCBwK6ihflS6MS3n&sz=w2000"

MARKERS = [
    "{{titulo}}",
    "{{artista}}",
    "{{subtitulo}}",
    "{{descripcion_corta}}",
    "{{fecha}}",
    "{{sinopsis}}",
    "{{propuesta_texto}}",
    "{{nota_direccion}}",
    "{{requerimientos_texto}}",
    "{{requerimientos_detalle}}",
    "{{contacto_texto}}",
    "{{cierre_texto}}",
    "{{email}}",
    "{{web}}",
    "{{redes}}",
]

def normalize_image(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return PLACEHOLDER
    if value.startswith("drive.google.com"):
        try:
            file_id = value.split("id=")[-1].split("&")[0]
            return f"https://drive.google.com/thumbnail?id={file_id}&sz=w2000"
        except Exception:
            return PLACEHOLDER
    return value

def render_docx(data: dict, out_path: Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Inter"
    style.font.size = Pt(11)

    doc.add_paragraph("{{titulo}}")
    doc.add_paragraph("{{artista}}")
    doc.add_paragraph("{{subtitulo}}")
    doc.add_paragraph("{{descripcion_corta}}")
    doc.add_paragraph("{{fecha}}")
    doc.add_paragraph("")
    doc.add_paragraph("Sinopsis")
    img = normalize_image(data.get("imagen_1") or data.get("IMAGEN_1") or "")
    if img:
        try:
            doc.add_picture(str(img), width=doc.sections[0].page_width * 0.6)
        except Exception:
            doc.add_paragraph("[Imagen no disponible]")
    doc.add_paragraph("{{sinopsis}}")
    doc.add_paragraph("")
    doc.add_paragraph("Propuesta artística")
    doc.add_paragraph("{{propuesta_texto}}")
    doc.add_paragraph("")
    doc.add_paragraph("{{nota_direccion}}")
    doc.add_paragraph("")
    doc.add_paragraph("Requerimientos técnicos")
    doc.add_paragraph("{{requerimientos_texto}}")
    doc.add_paragraph("{{requerimientos_detalle}}")
    doc.add_paragraph("")
    doc.add_paragraph("Contacto")
    doc.add_paragraph("{{contacto_texto}}")
    doc.add_paragraph("{{email}}")
    doc.add_paragraph("{{web}}")
    doc.add_paragraph("{{redes}}")
    doc.add_paragraph("")
    doc.add_paragraph("{{cierre_texto}}")

    out = out_path or (WORKSPACE / "output" / f"{data.get('titulo') or data.get('TITULO') or 'dossier'}.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out

if __name__ == "__main__":
    import json
    payload = Path.home() / "athamu_workspace" / "payload_flautisc0.json"
    data = json.loads(payload.read_text(encoding="utf-8"))
    print(render_docx(data, None))
