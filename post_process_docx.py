#!/usr/bin/env python3
"""Post-process ATHAMU DOCX after docxtpl render."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TOKENS = {
    "dark": {
        "bg": "#0D0D11",
        "card": "#16161F",
        "text": "#FFFFFF",
        "secondary": "#A0A0AB",
        "gold": "#FFB800",
        "red": "#E50914",
    },
    "light": {
        "bg": "#FFFFFF",
        "card": "#F3F3F3",
        "text": "#111111",
        "secondary": "#444444",
        "gold": "#D4AF37",
        "red": "#B91C1C",
    },
}


def hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def apply_style(doc: Document, variant: str):
    palette = TOKENS.get(variant.lower(), TOKENS["dark"])

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    body = doc.styles.get("Normal")
    if body:
        body.font.name = "Inter"
        body.font.size = Pt(11)
        body.font.color.rgb = RGBColor(*hex_to_rgb(palette["text"]))

    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = doc.styles.get(style_name)
        if style:
            style.font.name = "Montserrat"
            style.font.size = Pt(style_name == "Title" and 32 or 20)
            style.font.bold = True
            style.font.color.rgb = RGBColor(*hex_to_rgb(palette["gold"] if style_name == "Title" else palette["red"]))

    body_paragraphs = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text == "Sinopsis" or text == "Galería" or text == "Contacto" or text.startswith("Requerimientos") or text.startswith("Propuesta"):
            para.style = doc.styles.get("Heading 1", para.style)
            if para.style.name != "Heading 1":
                try:
                    run = para.runs[0] if para.runs else para.add_run(text)
                    run.font.name = "Montserrat"
                    run.font.size = Pt(20)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(*hex_to_rgb(palette["red"]))
                except Exception:
                    pass
        else:
            for run in para.runs:
                run.font.name = "Inter"
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(*hex_to_rgb(palette["secondary"]))
            body_paragraphs.append(para)


def save(doc: Document, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def post_process(docx_path: Path, variant: str = "cinematic-dark") -> Path:
    doc = Document(str(docx_path))
    try:
        apply_style(doc, variant)
    except Exception:
        pass
    out = docx_path.with_name(f"{docx_path.stem}_styled.docx")
    save(doc, out)
    return out


if __name__ == "__main__":
    import sys
    from pathlib import Path
    p = Path(sys.argv[1])
    v = sys.argv[2] if len(sys.argv) > 2 else "cinematic-dark"
    print(post_process(p, v))
