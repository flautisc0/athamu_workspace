#!/usr/bin/env python3
"""Generate ATHAMU DOCX templates for docxtpl."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

WORKSPACE = Path.home() / "athamu_workspace"
TOKENS_PATH = WORKSPACE / "tokens" / "atha-design-tokens.json"

def load_tokens():
    if TOKENS_PATH.exists():
        try:
            return json.loads(TOKENS_PATH.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}

def hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    return RGBColor(r, g, b)

def set_run_font(run, font_name, size_pt=None, color=None, bold=False, italic=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size_pt:
        run.font.size = Pt(size_pt)
    if isinstance(color, str):
        color = hex_to_rgb(color)
    if color is not None:
        run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic

def add_placeholder_paragraph(doc, text, style_name="Normal", font_name="Inter", size_pt=11.5, color_hex=None, bold=False, italic=False, align=None):
    p = doc.add_paragraph(style=style_name)
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, font_name, size_pt, color_hex, bold=bold)
    run.font.italic = italic
    return p

def apply_page_margins(doc, left_cm=2.54, top_cm=2.54, right_cm=2.54, bottom_cm=2.54):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.orientation = 0
    section.top_margin = Cm(top_cm)
    section.bottom_margin = Cm(bottom_cm)
    section.left_margin = Cm(left_cm)
    section.right_margin = Cm(right_cm)

def add_heading_style(doc, name="Heading 1", font_name="Montserrat", size_pt=25, color_hex="#FFFFFF"):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, 1)
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.font.size = Pt(size_pt)
    style.font.bold = True
    style.font.color.rgb = hex_to_rgb(color_hex)
    style.paragraph_format.space_after = Pt(10)
    style.paragraph_format.space_before = Pt(0)

def add_body_style(doc, name="Body", font_name="Inter", size_pt=11.5, color_hex="#A0A0AB"):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, 1)
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    style.font.size = Pt(size_pt)
    style.font.color.rgb = hex_to_rgb(color_hex)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(0)

def build_cinematic_dark(path: Path, tokens):
    doc = Document()
    apply_page_margins(doc, left_cm=1.6, top_cm=1.4, right_cm=1.6, bottom_cm=1.8)
    add_heading_style(doc, "CoverTitle", "Montserrat", 54, tokens.get("colors", {}).get("text-primary", "#FFFFFF"))
    add_heading_style(doc, "CoverSubTitle", "Montserrat", 17, tokens.get("colors", {}).get("accent-gold", "#FFB800"))
    add_heading_style(doc, "SectionTitle", "Montserrat", 25, tokens.get("colors", {}).get("text-primary", "#FFFFFF"))
    add_body_style(doc, "Body", "Inter", 11.5, tokens.get("colors", {}).get("text-secondary", "#A0A0AB"))
    add_body_style(doc, "CardText", "Inter", 11, "#A0A0AB")
    add_body_style(doc, "FooterText", "Inter", 9.5, "#A0A0AB")
    # Cover
    add_placeholder_paragraph(doc, "Dossier Artístico", "Normal", "Inter", 9.5, tokens.get("colors", {}).get("accent-red", "#E50914"), bold=True)
    add_placeholder_paragraph(doc, "{{titulo}}", "CoverTitle", "Montserrat", 54, tokens.get("colors", {}).get("text-primary", "#FFFFFF"), bold=True)
    add_placeholder_paragraph(doc, "{{subtitulo}}", "CoverSubTitle", "Montserrat", 17, tokens.get("colors", {}).get("accent-gold", "#FFB800"), bold=True)
    add_placeholder_paragraph(doc, "{{descripcion_corta}}", "Body", "Inter", 11.5, tokens.get("colors", {}).get("text-secondary", "#A0A0AB"))
    footer_p = add_placeholder_paragraph(doc, "ATHA Producciones — {{fecha}}", "FooterText", "Inter", 9.5, "#A0A0AB")
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_placeholder_paragraph(doc, "{{logo_atha}}", "FooterText", "Inter", 9.5, "#A0A0AB")
    # Sinopsis
    add_placeholder_paragraph(doc, "Sinopsis", "SectionTitle")
    add_placeholder_paragraph(doc, "{{sinopsis}}", "Body")
    add_placeholder_paragraph(doc, "[Image: {{imagen_1}}]", "Body", "Inter", 10, "#A0A0AB")
    # Propuesta
    add_placeholder_paragraph(doc, "Propuesta", "SectionTitle")
    add_placeholder_paragraph(doc, "{{propuesta_texto}}", "Body")
    add_placeholder_paragraph(doc, "{{nota_direccion}}", "CardText", "Inter", 11, hex_to_rgb("#A0A0AB"), italic=True)
    # Requerimientos
    add_placeholder_paragraph(doc, "Requerimientos", "SectionTitle")
    add_placeholder_paragraph(doc, "{{requerimientos_texto}}", "Body")
    add_placeholder_paragraph(doc, "{{requerimientos_detalle}}", "Body")
    # Contacto
    add_placeholder_paragraph(doc, "Contacto", "SectionTitle")
    add_placeholder_paragraph(doc, "{{contacto_texto}}", "Body")
    add_placeholder_paragraph(doc, "Email: {{email}}", "Body")
    add_placeholder_paragraph(doc, "Web: {{web}}", "Body")
    add_placeholder_paragraph(doc, "Redes: {{redes}}", "Body")
    # Cierre + footer logo placeholder
    add_placeholder_paragraph(doc, "{{cierre_texto}}", "Body")
    add_placeholder_paragraph(doc, "[LOGO PLACEHOLDER]", "FooterText")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path

def build_light(path: Path, tokens):
    doc = Document()
    apply_page_margins(doc, left_cm=1.6, top_cm=1.4, right_cm=1.6, bottom_cm=1.8)
    text_primary = tokens.get("colors", {}).get("text-primary", "#111111")
    text_secondary = tokens.get("colors", {}).get("text-secondary", "#444444")
    accent_gold = tokens.get("colors", {}).get("accent-gold", "#D4AF37")
    accent_red = tokens.get("colors", {}).get("accent-red", "#B91C1C")
    add_heading_style(doc, "CoverTitle", "Montserrat", 54, text_primary)
    add_heading_style(doc, "CoverSubTitle", "Montserrat", 17, accent_red)
    add_heading_style(doc, "SectionTitle", "Montserrat", 25, text_primary)
    add_body_style(doc, "Body", "Inter", 11.5, text_secondary)
    add_body_style(doc, "CardText", "Inter", 11, text_secondary)
    add_body_style(doc, "FooterText", "Inter", 9.5, text_secondary)
    add_placeholder_paragraph(doc, "Dossier Artístico", "Normal", "Inter", 9.5, accent_red, bold=True)
    add_placeholder_paragraph(doc, "{{titulo}}", "CoverTitle", "Montserrat", 54, text_primary, bold=True)
    add_placeholder_paragraph(doc, "{{subtitulo}}", "CoverSubTitle", "Montserrat", 17, accent_red, bold=True)
    add_placeholder_paragraph(doc, "{{descripcion_corta}}", "Body")
    footer_p = add_placeholder_paragraph(doc, "ATHA Producciones — {{fecha}}", "FooterText")
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_placeholder_paragraph(doc, "{{logo_atha}}", "FooterText", "Inter", 9.5, text_secondary)
    add_placeholder_paragraph(doc, "Sinopsis", "SectionTitle")
    add_placeholder_paragraph(doc, "{{sinopsis}}", "Body")
    add_placeholder_paragraph(doc, "[Image: {{imagen_1}}]", "Body", "Inter", 10, text_secondary)
    add_placeholder_paragraph(doc, "Propuesta", "SectionTitle")
    add_placeholder_paragraph(doc, "{{propuesta_texto}}", "Body")
    add_placeholder_paragraph(doc, "{{nota_direccion}}", "CardText", "Inter", 11, text_secondary, italic=True)
    add_placeholder_paragraph(doc, "Requerimientos", "SectionTitle")
    add_placeholder_paragraph(doc, "{{requerimientos_texto}}", "Body")
    add_placeholder_paragraph(doc, "{{requerimientos_detalle}}", "Body")
    add_placeholder_paragraph(doc, "Contacto", "SectionTitle")
    add_placeholder_paragraph(doc, "{{contacto_texto}}", "Body")
    add_placeholder_paragraph(doc, "Email: {{email}}", "Body")
    add_placeholder_paragraph(doc, "Web: {{web}}", "Body")
    add_placeholder_paragraph(doc, "Redes: {{redes}}", "Body")
    add_placeholder_paragraph(doc, "{{cierre_texto}}", "Body")
    add_placeholder_paragraph(doc, "[LOGO PLACEHOLDER]", "FooterText")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path

def main():
    tokens = load_tokens()
    print("Loaded tokens:", bool(tokens))
    dark_path = WORKSPACE / "templates" / "cinematic-dark" / "template.docx"
    light_path = WORKSPACE / "templates" / "light" / "template.docx"
    build_cinematic_dark(dark_path, tokens)
    build_light(light_path, tokens)
    print("Created:", dark_path)
    print("Created:", light_path)

if __name__ == "__main__":
    main()
